"""Present aggregate availability below newly generated report metric tables."""

import re
from collections.abc import Mapping
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

_COVERAGE_SUFFIX = " coverage"
_METRIC_LABELS = {"ap50": "AP50", "ap75": "AP75", "ap50_95": "AP50-95"}


def compact_report_counts(
    path: Path,
    sheet_names: tuple[str, ...],
    *,
    translations: Mapping[str, str] | None = None,
) -> None:
    """Move upstream mean-row counts without recalculating their populations.

    Call only on fresh report-generator outputs, before identity annotation.
    Their metric headers occupy row 2 and generated coverage columns are trailing.
    """
    workbook = load_workbook(path)
    try:
        for sheet in workbook:
            if sheet.title in sheet_names:
                _compact_sheet(sheet, translations or {})
        workbook.save(path)
    finally:
        workbook.close()


def _compact_sheet(sheet: Any, translations: Mapping[str, str]) -> None:
    """Keep the untyped openpyxl worksheet API inside the rendering boundary."""
    columns = _summary_columns(sheet, translations)
    if not columns:
        return
    # Column deletion does not rewrite formulas or shift column dimensions.
    # The upstream builders produce static tables with appended coverage columns.
    if [column for column, _, _ in columns] != list(range(columns[0][0], sheet.max_column + 1)):
        raise ValueError(f"Report {sheet.title!r} has non-trailing coverage columns")
    for column, _, _ in reversed(columns):
        sheet.delete_cols(column)
        letter = get_column_letter(column)
        if letter in sheet.column_dimensions:
            del sheet.column_dimensions[letter]

    start = sheet.max_row + 2
    title = sheet.cell(start, 1, "Valid class counts (valid / eligible)")
    title.font = Font(name="Calibri", bold=True)
    title.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[start].height = 32
    for row, (_, label, value) in enumerate(columns, start + 1):
        name = sheet.cell(row, 1, label)
        name.data_type = "s"
        name.alignment = Alignment(wrap_text=True, vertical="center")
        name.font = copy(title.font)
        count = sheet.cell(row, 2, value)
        count.data_type = "s"
        count.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[row].height = 30
    sheet.column_dimensions["A"].width = max(sheet.column_dimensions["A"].width, 36)


def _summary_columns(sheet: Any, translations: Mapping[str, str]) -> list[tuple[int, str, str]]:
    mean_rows = [cell.row for cell in sheet["A"] if cell.value == "Среднее"]
    if len(mean_rows) != 1:
        raise ValueError(f"Report {sheet.title!r} must have exactly one mean row")
    headers = {str(cell.value) for cell in sheet[2]}
    source_headers = {
        display: source
        for source, display in translations.items()
        if source.endswith(_COVERAGE_SUFFIX)
        and translations.get(
            source.removesuffix(_COVERAGE_SUFFIX), source.removesuffix(_COVERAGE_SUFFIX)
        )
        in headers
    }
    columns: list[tuple[int, str, str]] = []
    for cell in sheet[2]:
        # Only generated aggregate cells contain valid/eligible strings. Ordinary
        # means are numbers or NA, even when their custom headings end in coverage.
        value = sheet.cell(mean_rows[0], cell.column).value
        if not isinstance(value, str) or re.fullmatch(r"\d+/\d+", value) is None:
            continue
        header = source_headers.get(str(cell.value), str(cell.value))
        if not header.endswith(_COVERAGE_SUFFIX):
            raise ValueError(f"Report {sheet.title!r} has an unknown class-count header {header!r}")
        metric = header.removesuffix(_COVERAGE_SUFFIX)
        label = _METRIC_LABELS.get(metric, metric)
        columns.append((cell.column, f"{label}-valid-class-count", value))
    return columns

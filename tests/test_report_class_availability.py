"""Real workbook builders retain metrics across changing model vocabularies."""

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils import column_index_from_string  # type: ignore[import-untyped]
from report_generator.config import Config

from clearml_yolo.adapters.clearml.session import ClearMLConfig
from clearml_yolo.adapters.reporting.reports import build_reports
from clearml_yolo.adapters.reporting.workbook_identity import (
    deannotated_workbook,
    workbook_identities,
)
from clearml_yolo.application.ports import WorkflowDependencies
from clearml_yolo.application.use_cases.compare import MANIFEST_NAME, ComparisonManifest
from clearml_yolo.application.use_cases.report import report
from workflow_dependencies import patch_workflow
from workflow_dependencies import workflow_dependencies as workflow_dependencies  # noqa: PLC0414


@pytest.mark.parametrize("shared", [True, False])
def test_report_preserves_one_sided_metrics_with_real_builders(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    shared: bool,
    workflow_dependencies: WorkflowDependencies,
) -> None:
    patch_workflow(
        monkeypatch,
        workflow_dependencies,
        "clearml_yolo.application.use_cases.report.init_task",
        lambda *_a, **_kw: None,
    )
    comparison = tmp_path / "comparison"
    comparison.mkdir()
    for model, unique, values in (
        ("candidate", "new", [0.8, 0.0]),
        ("baseline", "deleted", [0.6, 0.4]),
    ):
        labels = ["shared", unique] if shared else [unique]
        metrics = values if shared else values[1:]
        pd.DataFrame(
            {
                "Класс": labels,
                "Количество примеров train": [30] * len(labels),
                "f1_score": metrics,
                "ap50": metrics,
                "perebrak": [0.1] * len(labels),
                "nedobrak": [0.1] * len(labels),
            }
        ).to_excel(comparison / f"{model}.xlsx", index=False)
    manifest = ComparisonManifest(
        split="test",
        candidate_dashboard="candidate.xlsx",
        baseline_dashboard="baseline.xlsx",
        candidate_predictions="candidate.csv",
        baseline_predictions="baseline.csv",
        statistical_workbook="comparison.xlsx",
    )
    (comparison / MANIFEST_NAME).write_text(manifest.model_dump_json(), encoding="utf-8")
    result = report(
        comparison,
        tmp_path / "reports",
        ClearMLConfig(),
        baseline_label="fixture baseline",
        candidate_label="fixture candidate",
        deps=workflow_dependencies,
    )
    config = Config.load()
    for business, path in (
        (False, result.dev_reports["test"]),
        (True, result.business_reports["test"]),
    ):
        identities = workbook_identities(path)
        assert identities["baseline"].model_name == "fixture baseline"
        assert identities["candidate"].model_name == "fixture candidate"
        annotated = load_workbook(path)
        for sheet in annotated:
            assert sheet.page_setup.fitToWidth == 1
            assert sheet.print_title_rows.startswith("$1:")
        annotated.close()
        with deannotated_workbook(path) as original:
            workbook = load_workbook(original)
        metric = (
            config.business.column_translations.get("f1_score", "f1_score")
            if business
            else "f1_score"
        )
        for sheet_name, available, absent, expected, mean in (
            (config.sheet_names.model1, "new", "deleted", 0.0, 0.4 if shared else 0.0),
            (config.sheet_names.model2, "deleted", "new", 0.4, 0.5 if shared else 0.4),
        ):
            sheet = workbook[sheet_name]
            headers = [cell.value for cell in sheet[2]]
            assert not any(str(header).endswith(" coverage") for header in headers)
            column = headers.index(metric) + 1
            rows = {sheet.cell(row, 1).value: row for row in range(3, sheet.max_row + 1)}
            assert sheet.cell(rows["AP50-valid-class-count"], 2).value == (
                "2/2" if shared else "1/1"
            )
            assert {"new", "deleted"} <= rows.keys()
            assert sheet.cell(rows[available], column).value == pytest.approx(expected)
            assert sheet.cell(rows[absent], column).value == "NA"
            assert sheet.cell(rows["Среднее"], column).value == pytest.approx(mean)
        sheet = workbook[config.sheet_names.comparison]
        headers = [cell.value for cell in sheet[2]]
        column = headers.index(metric) + 1
        rows = {sheet.cell(row, 1).value: row for row in range(3, sheet.max_row + 1)}
        assert sheet.cell(rows["AP50-valid-class-count"], 2).value == ("1/1" if shared else "0/0")
        for label in ("new", "deleted"):
            cell = sheet.cell(rows[label], column)
            assert cell.value == "NA"
            assert cell.fill.patternType is None
        workbook.close()


@pytest.mark.parametrize(
    ("ap50", "train_counts", "expected"),
    [
        ([0.0, 0.8], [30, 30], "2/2"),
        ([0.0, None], [30, 30], "1/2"),
        ([None, None], [30, 30], "0/2"),
        ([0.0, 0.8], [0, 0], "0/0"),
    ],
)
def test_report_count_summary_preserves_values_and_custom_headers(
    tmp_path: Path,
    ap50: list[float | None],
    train_counts: list[int],
    expected: str,
) -> None:
    source = tmp_path / "source.xlsx"
    pd.DataFrame(
        {
            "Класс": ["A", "B"],
            "Количество примеров train": train_counts,
            "ap50": ap50,
            "f1_score": [0.8, 0.7],
            "perebrak": [0.1, 0.2],
            "nedobrak": [0.2, 0.1],
            "Availability custom": [0.4, 0.6],
        }
    ).to_excel(source, index=False)
    original = source.read_bytes()
    settings = tmp_path / "report.yaml"
    settings.write_text(
        "business:\n  column_translations:\n"
        "    ap50: AP custom\n    ap50 coverage: AP availability\n"
        "    absent coverage: Availability custom\n"
        "    unused: f1_score coverage\n",
        encoding="utf-8",
    )
    dev, business = tmp_path / "dev.xlsx", tmp_path / "business.xlsx"
    build_reports(source, source, dev, business, settings)
    assert source.read_bytes() == original
    config = Config.load(settings)
    for path in (dev, business):
        workbook = load_workbook(path)
        for sheet_name in (
            config.sheet_names.model1,
            config.sheet_names.model2,
            config.sheet_names.comparison,
        ):
            sheet = workbook[sheet_name]
            headers = [cell.value for cell in sheet[2]]
            assert "AP availability" not in headers
            custom_present = expected != "0/0" or sheet_name != config.sheet_names.comparison
            assert ("Availability custom" in headers) == custom_present
            assert not any(str(header).endswith(" coverage") for header in headers)
            metric = "AP custom" if path == business else "ap50"
            column = headers.index(metric) + 1
            rows = {sheet.cell(row, 1).value: row for row in range(3, sheet.max_row + 1)}
            assert sheet.cell(rows["AP50-valid-class-count"], 2).value == expected
            assert sheet.cell(rows["f1_score-valid-class-count"], 2).value == (
                "0/0" if expected == "0/0" else "2/2"
            )
            if custom_present:
                assert sheet.cell(rows["Availability custom-valid-class-count"], 2).value == (
                    "0/0" if expected == "0/0" else "2/2"
                )
            assert rows["AP50-valid-class-count"] > rows["Среднее"]
            assert sheet.freeze_panes == "B3"
            assert all(
                column_index_from_string(key) <= len(headers) for key in sheet.column_dimensions
            )
            if expected != "0/0":
                assert sheet.cell(rows["A"], column).value == ("NA" if ap50[0] is None else ap50[0])
                if ap50[0] is not None:
                    assert sheet.cell(rows["A"], column).number_format == "0.0000"
            if expected.startswith("0/"):
                assert sheet.cell(rows["Среднее"], column).value == "NA"
        workbook.close()

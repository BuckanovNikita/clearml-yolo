"""Evaluation computation returns neutral values without producing artifacts."""

from pathlib import Path
from typing import Any

import pandas as pd
import pytest

from clearml_yolo.adapters.evaluation import scoring


def test_computation_is_available_without_report_output_paths(tmp_path: Path) -> None:
    assert hasattr(scoring, "compute_evaluation"), "Scoring must expose computation separately"
    gt = pd.DataFrame(
        [("a", "cat", 0, 0, 10, 10)],
        columns=[
            "image_name",
            "instance_label",
            "bbox_x_tl",
            "bbox_y_tl",
            "bbox_x_br",
            "bbox_y_br",
        ],
    ).assign(split="test")
    predictions = gt.drop(columns="split").assign(confidence=0.7)
    computed = scoring.compute_evaluation(
        gt,
        predictions,
        predictions,
        split="test",
        classes=["cat"],
        thresholds={"cat": 0.7},
        required_classes=["cat"],
        iou_threshold=0.5,
        matching_strategy="iou_prior",
        ap_method="interp",
        skip_cohen_kappa=True,
    )
    assert computed.metrics["cat"].tp == 1
    assert computed.metrics["cat"].confidence == 0.7
    assert computed.metrics["cat"].__class__.__module__ == "clearml_yolo.core.evaluation.models"
    assert computed.confusion_matrix.counts == [[1, 0], [0, 0]]
    assert list(tmp_path.iterdir()) == []


def test_rendered_metrics_preserve_every_backend_value(tmp_path: Path) -> None:
    from digital_metrics.engines import compute_metrics_from_matches
    from digital_metrics.matching import match_boxes
    from digital_metrics.scoring import compute_map, slice_by_conf

    from clearml_yolo.adapters.reporting.evaluation import render_evaluation, summarize_evaluation
    from clearml_yolo.core.evaluation.models import EvaluatedSplit

    columns = ["image_name", "instance_label", "bbox_x_tl", "bbox_y_tl", "bbox_x_br", "bbox_y_br"]
    gt = pd.DataFrame([("a", "cat", 0, 0, 10, 10)], columns=columns).assign(split="test")
    predictions = gt.drop(columns="split").assign(confidence=0.700000000000003)
    threshold = 0.700000000000003
    computed = scoring.compute_evaluation(
        gt,
        predictions,
        predictions,
        split="test",
        classes=["cat"],
        thresholds={"cat": threshold},
        required_classes=["cat"],
        iou_threshold=0.5,
        matching_strategy="iou_prior",
        ap_method="interp",
        skip_cohen_kappa=True,
    )
    upstream_matches = match_boxes(
        gt, predictions, 0.5, strategy="iou_prior", split_image_names=["a"]
    )
    upstream_metrics = compute_metrics_from_matches(
        slice_by_conf(upstream_matches, ["cat"], {"cat": threshold}),
        ["cat"],
        {"cat": threshold},
    )
    compute_map(gt, predictions, upstream_metrics, ["a"], method="interp", strategy="iou_prior")
    assert computed.metrics["cat"].model_dump() == upstream_metrics["cat"].model_dump()
    artifacts = render_evaluation(computed, output_dir=tmp_path, suffix="test")
    evaluated = EvaluatedSplit.from_parts(computed, artifacts)
    assert evaluated.metrics is computed.metrics
    assert evaluated.result_rows is computed.result_rows
    assert evaluated.dashboard_path.is_file()
    assert evaluated.dtrk_dashboard_path.is_file()
    assert evaluated.confusion_matrix_path.is_file()
    assert all(path.is_file() for path in evaluated.plot_paths.values())
    per_class, summary = summarize_evaluation(evaluated.metrics)
    assert per_class.loc["cat", "confidence"] == threshold
    assert summary["mean_precision"] == summary["mean_recall"] == 1.0


def test_source_row_assembly_receives_project_owned_match_values(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from clearml_yolo.core.evaluation.result_rows import build_result_rows

    original = build_result_rows
    observed_modules: list[str] = []

    def observed(*args: Any, **kwargs: Any) -> pd.DataFrame:
        for name in ("matches_pre_threshold", "matches_post_threshold"):
            observed_modules.extend(
                match.__class__.__module__ for matches in kwargs[name].values() for match in matches
            )
        return original(*args, **kwargs)

    columns = ["image_name", "instance_label", "bbox_x_tl", "bbox_y_tl", "bbox_x_br", "bbox_y_br"]
    gt = pd.DataFrame([("a", "cat", 0, 0, 10, 10)], columns=columns).assign(split="test")
    predictions = gt.drop(columns="split").assign(confidence=0.7)
    monkeypatch.setattr(scoring, "build_result_rows", observed)
    scoring.compute_evaluation(
        gt,
        predictions,
        predictions,
        split="test",
        classes=["cat"],
        thresholds={"cat": 0.7},
        required_classes=["cat"],
        iou_threshold=0.5,
        matching_strategy="iou_prior",
        ap_method="interp",
        skip_cohen_kappa=True,
    )
    assert observed_modules == ["clearml_yolo.core.evaluation.models"] * 2


def test_opaque_evaluation_classes_are_readable_without_changing_machine_values(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from zipfile import ZipFile

    from digital_metrics.reporting import render_evaluation as shared
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    from clearml_yolo.adapters.reporting import evaluation
    from clearml_yolo.adapters.reporting.evaluation_workbook import write_candidate_workbook
    from clearml_yolo.adapters.reporting.workbook_identity import read_dashboard

    names = ["0123456789abcdef" * 2, "fedcba9876543210" * 2]
    columns = ["image_name", "instance_label", "bbox_x_tl", "bbox_y_tl", "bbox_x_br", "bbox_y_br"]
    truth = pd.DataFrame(
        [(name, name, 0, 0, 10, 10) for name in names],
        columns=columns,
    ).assign(split="test")
    predictions = truth.drop(columns="split").assign(confidence=0.75)
    computed = scoring.compute_evaluation(
        truth,
        predictions,
        predictions,
        split="test",
        classes=names,
        thresholds=dict.fromkeys(names, 0.75),
        required_classes=names,
        iou_threshold=0.5,
        matching_strategy="iou_prior",
        ap_method="interp",
        skip_cohen_kappa=True,
    )
    original_truth = computed.ground_truth.copy(deep=True)
    plots: list[Any] = []

    def render(*args: Any, **kwargs: Any) -> Any:
        result = shared(*args, **kwargs)
        plots.extend(result.plots.values())
        return result

    monkeypatch.setattr(evaluation, "render_shared_evaluation", render)
    artifacts = evaluation.render_evaluation(computed, output_dir=tmp_path, suffix="test")
    for plot in plots:
        labels = [label.get_text() for label in plot.axes.get_xticklabels()]
        assert len(set(labels)) == 2
        assert all(name not in labels for name in names)
    for path in (
        artifacts.dashboard_path,
        artifacts.dtrk_dashboard_path,
        artifacts.confusion_matrix_path,
    ):
        with ZipFile(path) as archive:
            assert not any(name.startswith("xl/media/") for name in archive.namelist())
        workbook = load_workbook(path)
        visible = " ".join(str(cell.value) for sheet in workbook for row in sheet for cell in row)
        assert all(name not in visible for name in names)
        workbook.close()
        assert set(names) <= set(read_dashboard(path, index_col=0).index)
    assert list(computed.metrics) == names
    assert list(artifacts.dashboard.index) == names
    assert artifacts.dashboard.loc[names[0], "confidence"] == 0.75
    pd.testing.assert_frame_equal(computed.ground_truth, original_truth)
    candidate = tmp_path / "candidate.xlsx"
    write_candidate_workbook(candidate, artifacts.dashboard, {"mean_precision": 1.0})
    assert read_dashboard(candidate, index_col=0).index.tolist() == names
    workbook = load_workbook(candidate)
    assert all(
        name not in str(cell.value)
        for sheet in workbook
        for row in sheet
        for cell in row
        for name in names
    )
    workbook.close()


def test_geometry_filter_validates_raw_evidence_before_dropping_rows(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from clearml_yolo.core.validation import validate_dataframe

    observed: list[tuple[str, int]] = []

    def validate(frame: pd.DataFrame, stage: str) -> pd.DataFrame:
        observed.append((stage, len(frame)))
        return validate_dataframe(frame, stage)

    monkeypatch.setattr(scoring, "validate_dataframe", validate)
    columns = ["image_name", "instance_label", "bbox_x_tl", "bbox_y_tl", "bbox_x_br", "bbox_y_br"]
    raw = pd.DataFrame(
        [
            ("a", "cat", 0, 0, 10, 10),
            ("a", "cat", 0, 0, 0, 10),
        ],
        columns=columns,
    ).assign(confidence=0.7)
    filtered = scoring.filter_invalid_prediction_boxes(raw)
    assert observed == [("raw_predictions", 2)]
    assert len(raw) == 2
    assert len(filtered) == 1

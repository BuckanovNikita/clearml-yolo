"""Real exported evaluation artifacts retain the evaluated model's identity."""

from pathlib import Path
from typing import Any

import pandas as pd
import pytest
from matplotlib.figure import Figure
from openpyxl import load_workbook  # type: ignore[import-untyped]
from PIL import Image

from clearml_yolo.adapters.clearml.report import (
    report_confusion_matrices,
    report_pr_curves,
    report_table,
)
from clearml_yolo.application.evaluation import evaluate_split
from clearml_yolo.application.ports import WorkflowDependencies
from clearml_yolo.core.evaluation.payload import EvaluationPayload
from clearml_yolo.core.evaluation.schema import ConfusionMatrixPayload, PRCurve, ResultContext
from clearml_yolo.core.identity import ModelIdentity
from workflow_dependencies import patch_workflow
from workflow_dependencies import (
    workflow_dependencies as workflow_dependencies,  # noqa: PLC0414 - fixture export
)


class RecordingTask:
    def __init__(self) -> None:
        self.plots: list[dict[str, Any]] = []
        self.tables: list[dict[str, Any]] = []

    def get_logger(self) -> "RecordingTask":
        return self

    def report_plotly(self, **kwargs: Any) -> None:
        self.plots.append(kwargs)

    def report_table(self, **kwargs: Any) -> None:
        self.tables.append(kwargs)


@pytest.mark.parametrize("training_id", [None, "0123456789abcdef0123456789abcdef"])
def test_real_exports_and_payload_retain_identity(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    training_id: str | None,
    workflow_dependencies: WorkflowDependencies,
) -> None:
    name = "=SUM(A1:A2) <b>猫</b> " + "long model " * 20
    identity = ModelIdentity(model_name=name, training_task_id=training_id)
    columns = ["image_name", "instance_label", "bbox_x_tl", "bbox_y_tl", "bbox_x_br", "bbox_y_br"]
    truth = pd.DataFrame([("a", "cat", 0, 0, 10, 10)], columns=columns).assign(split="test")
    predictions = pd.DataFrame([("a", "cat", 0, 0, 10, 10, 0.9)], columns=[*columns, "confidence"])
    captions: list[str] = []
    original = Figure.savefig

    def observe_export(figure: Figure, *args: Any, **kwargs: Any) -> None:
        captions.extend(text.get_text() for text in figure.texts)
        figure.canvas.draw()
        for text in figure.texts:
            assert text.get_window_extent().ymin > figure.axes[0].title.get_window_extent().ymax
        original(figure, *args, **kwargs)

    patch_workflow(monkeypatch, workflow_dependencies, Figure, "savefig", observe_export)
    evaluated = evaluate_split(
        truth,
        predictions,
        predictions,
        split="test",
        classes=["cat"],
        thresholds={"cat": 0.5},
        required_classes=None,
        iou_threshold=0.5,
        matching_strategy="iou_prior",
        ap_method="interp",
        skip_cohen_kappa=True,
        output_dir=tmp_path,
        suffix="test",
        model_identity=identity,
        deps=workflow_dependencies,
    )
    assert evaluated.model_identity == identity
    assert (
        EvaluationPayload.model_validate_json(
            evaluated.evaluation_payload.model_dump_json(),
        ).model_identity
        == identity
    )
    assert len(captions) == 4
    for caption in captions:
        assert "Training task: " + (training_id or "unavailable") in caption
        assert name.split(maxsplit=1)[0] in caption
    for plot in evaluated.plot_paths.values():
        with Image.open(plot) as image:
            image.verify()
            assert image.info["Model identity"] == (
                f"Model: {name}\nTraining task: {training_id or 'unavailable'}"
            )
    for path in (
        evaluated.dashboard_path,
        evaluated.dtrk_dashboard_path,
        evaluated.confusion_matrix_path,
    ):
        workbook = load_workbook(path)
        values = [cell for sheet in workbook for row in sheet.iter_rows() for cell in row]
        assert name in "".join(str(cell.value) for cell in values if cell.value is not None)
        assert all(cell.data_type != "f" for cell in values)
        assert any((training_id or "unavailable") in str(cell.value) for cell in values)
        workbook.close()


def test_plotly_captions_are_escaped_without_internal_ids() -> None:
    identity = ModelIdentity(model_name="=1 <b>猫 & model</b>", training_task_id="full-task-id")
    context = ResultContext(
        context_id="candidate", model_id="weights", split="test", model_identity=identity
    )
    old_context = context.model_copy(update={"model_identity": None})
    task, old_task = RecordingTask(), RecordingTask()
    matrix = ConfusionMatrixPayload(labels=["cat", "background"], counts=[[1, 0], [0, 0]])
    curve = PRCurve(
        class_name="cat",
        recall=[1],
        precision=[1],
        confidence=[0.9],
        tp=[1],
        fp=[0],
        gt_count=1,
        ap50=1,
        integration_method="interp",
    )
    for target, scope in ((task, context), (old_task, old_context)):
        report_confusion_matrices(target, scope, matrix)
        report_pr_curves(target, scope, [curve])
    for plot, old_plot in zip(task.plots, old_task.plots, strict=True):
        caption = plot["figure"]["layout"]["annotations"][0]["text"]
        assert "=1 &lt;b&gt;猫 &amp; model&lt;/b&gt;" in caption
        assert "full-task-id" not in caption
        assert plot["series"] == "=1 <b>猫 & model</b> · test"
        assert old_plot["series"] == "Current model · test"
        assert plot["figure"]["data"] == old_plot["figure"]["data"]


def test_display_table_identity_preserves_source_frame() -> None:
    task = RecordingTask()
    frame = pd.DataFrame({"precision": [1.0]}, index=["cat"])
    identity = ModelIdentity(model_name="local checkpoint")
    report_table(task, "metrics", "test", frame, identities={"model": identity})
    displayed = task.tables[0]["table_plot"]
    assert displayed["model_model_name"].tolist() == ["local checkpoint"]
    assert displayed["model_training_task_id"].tolist() == ["unavailable"]
    assert frame.columns.tolist() == ["precision"]


def test_empty_comparison_table_keeps_both_identities_visible() -> None:
    task = RecordingTask()
    frame = pd.DataFrame(columns=["class_name", "precision_delta"])
    identities = {
        "baseline": ModelIdentity(model_name="Old <model>", training_task_id="full-old-id"),
        "candidate": ModelIdentity(model_name="New 模型", training_task_id="full-new-id"),
    }
    report_table(task, "degraded_classes", "test", frame, identities=identities)
    reported = task.tables[0]
    assert reported["table_plot"].empty
    assert frame.columns.tolist() == ["class_name", "precision_delta"]
    caption = reported["extra_layout"]["annotations"][0]["text"]
    assert "baseline: Old &lt;model&gt;" in caption
    assert "candidate: New 模型" in caption
    assert "full-old-id" not in caption
    assert "full-new-id" not in caption
    assert reported["title"] == "degraded_classes"
    assert reported["series"] == "test"

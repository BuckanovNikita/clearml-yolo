"""Source identity takes precedence and legacy inputs require explicit labels."""

from pathlib import Path
from typing import Any

import pandas as pd
import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]

from clearml_yolo.clearml_results import write_prediction_provenance
from clearml_yolo.clearml_session import ClearMLConfig
from clearml_yolo.model_identity import (
    ModelIdentity,
    checkpoint_sha256,
    require_model_identity,
    write_checkpoint_identity,
)
from clearml_yolo.publishing.models import FiftyOneConfig
from clearml_yolo.tasks import predict as predict_module
from clearml_yolo.tasks.compare import ModelRef, _resolve_model
from clearml_yolo.tasks.metrics import EvaluationConfig, compute_metrics
from clearml_yolo.tasks.report import report
from clearml_yolo.workbook_identity import annotate_workbook, workbook_identities
from test_metrics import _write_inputs
from test_predict import _predict
from test_predict import checkpoint_recording as checkpoint_recording  # noqa: PLC0414
from test_predict import published as published  # noqa: PLC0414
from test_report import _comparison_dir
from test_report import report_generator as report_generator  # noqa: PLC0414


@pytest.mark.parametrize("label", [None, "", "  "])
def test_unknown_model_requires_nonempty_label(label: str | None) -> None:
    with pytest.raises(ValueError, match="model_label"):
        require_model_identity(None, label)


def test_custom_label_does_not_invent_training_task() -> None:
    identity = require_model_identity(None, "legacy detector")
    assert identity.model_name == "legacy detector"
    assert identity.training_task_id is None
    assert identity.caption.endswith("Training task: unavailable")


def test_metrics_legacy_csv_requires_label_before_reading(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr("clearml_yolo.tasks.metrics.init_task", lambda *_args, **_kwargs: None)
    with pytest.raises(ValueError, match="model_label"):
        compute_metrics(
            tmp_path / "legacy.csv",
            tmp_path / "truth.csv",
            tmp_path / "metrics",
            ClearMLConfig(),
            EvaluationConfig(),
        )


def test_predict_keeps_recovered_training_identity_over_custom_label(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    checkpoint_recording: Any,
    published: dict[str, pd.DataFrame],
) -> None:
    checkpoint = tmp_path / "best.pt"
    checkpoint.write_bytes(b"verified checkpoint")
    identity = ModelIdentity(
        model_name="detector-calm-otter",
        training_task_id="a" * 32,
        checkpoint_sha256=checkpoint_sha256(checkpoint),
        model_id="trained-model",
    )
    monkeypatch.setattr(
        predict_module,
        "resolve_weights_with_identity",
        lambda _: (checkpoint, identity),
    )
    checkpoint_recording({"imgsz": 64})
    result = _predict(tmp_path, 64)
    assert result.model_identity == identity


def test_predict_unknown_checkpoint_requires_label(
    tmp_path: Path,
    checkpoint_recording: Any,
    published: dict[str, pd.DataFrame],
) -> None:
    checkpoint_recording({"imgsz": 64})
    with pytest.raises(ValueError, match="model_label"):
        _predict(tmp_path, 64, model_label=None)


def test_report_retains_distinct_source_tasks_and_visible_names(
    tmp_path: Path,
    report_generator: list[tuple[Path, Path]],
) -> None:
    comparison, candidate, baseline = _comparison_dir(tmp_path)
    identities = {
        "baseline": ModelIdentity(model_name="baseline-trained", training_task_id="a" * 32),
        "candidate": ModelIdentity(model_name="candidate-trained", training_task_id="b" * 32),
    }
    annotate_workbook(candidate, {"model": identities["candidate"]})
    annotate_workbook(baseline, {"model": identities["baseline"]})
    result = report(
        comparison,
        tmp_path / "reports",
        ClearMLConfig(),
        baseline_label="wrong baseline",
        candidate_label="wrong candidate",
    )
    for path in [*result.dev_reports.values(), *result.business_reports.values()]:
        assert workbook_identities(path) == identities
        workbook = load_workbook(path)
        visible = "\n".join(
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
        for identity in identities.values():
            assert identity.model_name in visible
            assert identity.training_task_id is not None
            assert identity.training_task_id in visible
        workbook.close()


def test_legacy_report_workbooks_require_explicit_labels(
    tmp_path: Path,
    report_generator: list[tuple[Path, Path]],
) -> None:
    comparison, _, _ = _comparison_dir(tmp_path)
    with pytest.raises(ValueError, match="model_label"):
        report(comparison, tmp_path / "reports", ClearMLConfig())


def test_local_comparison_without_provenance_requires_label(tmp_path: Path) -> None:
    checkpoint = tmp_path / "legacy.pt"
    checkpoint.write_bytes(b"legacy")
    with pytest.raises(ValueError, match="label"):
        _resolve_model(
            ModelRef(source="local", weights=checkpoint, thresholds={"cat": 0.5}),
            "project",
            exclude_task_id=None,
            automatic_absence_is_skip=False,
        )


def test_local_comparison_preserves_checkpoint_source_over_label(tmp_path: Path) -> None:
    checkpoint = tmp_path / "best.pt"
    checkpoint.write_bytes(b"known training checkpoint")
    identity = ModelIdentity(
        model_name="source-trained",
        training_task_id="a" * 32,
        checkpoint_sha256=checkpoint_sha256(checkpoint),
    )
    write_checkpoint_identity(checkpoint, identity)
    resolved = _resolve_model(
        ModelRef(
            source="local",
            weights=checkpoint,
            thresholds={"cat": 0.5},
            label="ignored label",
        ),
        "project",
        exclude_task_id=None,
        automatic_absence_is_skip=False,
    )
    assert resolved.identity == identity


def test_metrics_retains_prediction_source_identity_in_dashboard(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    predictions, truth = _write_inputs(tmp_path)
    checkpoint = tmp_path / "best.pt"
    checkpoint.write_bytes(b"training checkpoint")
    identity = ModelIdentity(
        model_name="trained-source",
        training_task_id="a" * 32,
        checkpoint_sha256=checkpoint_sha256(checkpoint),
    )
    write_prediction_provenance(predictions, checkpoint, identity)
    monkeypatch.setattr("clearml_yolo.tasks.metrics.init_task", lambda *_args, **_kwargs: None)
    result = compute_metrics(
        predictions,
        truth,
        tmp_path / "metrics",
        ClearMLConfig(),
        EvaluationConfig(),
        splits=["test"],
        model_label="ignored custom label",
        fiftyone=FiftyOneConfig(enabled=False),
    )
    assert workbook_identities(result.dashboards["test"]) == {"model": identity}

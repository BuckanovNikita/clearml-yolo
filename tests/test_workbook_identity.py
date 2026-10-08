"""Workbook identities preserve source data while repeating in print titles."""

from pathlib import Path
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font  # type: ignore[import-untyped]
from openpyxl.worksheet.pagebreak import Break  # type: ignore[import-untyped]
from openpyxl.worksheet.table import Table  # type: ignore[import-untyped]

from clearml_yolo.adapters.reporting.workbook_identity import (
    annotate_workbook,
    deannotated_workbook,
    read_dashboard,
    workbook_identities,
)
from clearml_yolo.core.identity import ModelIdentity


def _complex_workbook(path: Path) -> None:
    workbook = Workbook()
    main = workbook.active
    main.title = "Main sheet"
    main.append(["Class", "Value"])
    main.append(["cat", 2])
    main.append(["total", "=SUM($B$2:B2)"])
    main["A1"].font = Font(bold=True, color="FF0000")
    main.merge_cells("D2:E2")
    main.add_table(Table(displayName="Results", ref="A1:B2"))
    main.freeze_panes = "B2"
    main.print_title_rows = "1:1"
    main.print_title_cols = "A:A"
    main.print_area = "A1:E3"
    main.row_breaks.append(Break(id=2))
    other = workbook.create_sheet("Other")
    other["A1"] = "='Main sheet'!$B$2+1"
    other.print_title_rows = "1:2"
    workbook.save(path)


def test_banner_all_sheets_preserves_references_and_print_structure(tmp_path: Path) -> None:
    path = tmp_path / "workbook.xlsx"
    _complex_workbook(path)
    original = path.read_bytes()
    identity = ModelIdentity(
        model_name="=Unicode 模型 " + "long name " * 40, training_task_id="full-training-task-id"
    )
    annotate_workbook(path, {"Candidate": identity})
    annotated = load_workbook(path)
    for sheet in annotated:
        banner_end = int(sheet.print_title_rows.split(":")[-1].replace("$", "")) - (
            1 if sheet.title == "Main sheet" else 2
        )
        banner = "".join(str(sheet.cell(row, 1).value or "") for row in range(1, banner_end + 1))
        assert identity.model_name in banner
        assert identity.training_task_id is not None
        assert identity.training_task_id in banner
        assert all(sheet.cell(row, 1).data_type == "s" for row in range(1, banner_end + 1))
        assert sheet.print_title_rows.startswith("$1:")
        assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
        assert sheet.page_setup.fitToWidth == 1
        assert all(sheet.row_dimensions[row].height <= 48 for row in range(1, banner_end + 1))
    shift = (
        next(cell.row for row in annotated["Main sheet"] for cell in row if cell.value == "Class")
        - 1
    )
    main = annotated["Main sheet"]
    assert main.cell(3 + shift, 2).value == f"=SUM($B${2 + shift}:B{2 + shift})"
    assert annotated["Other"].cell(1 + shift, 1).value == f"='Main sheet'!$B${2 + shift}+1"
    assert main.cell(1 + shift, 1).font.color.rgb == "00FF0000"
    assert f"D{2 + shift}:E{2 + shift}" in main.merged_cells
    assert main.tables["Results"].ref == f"A{1 + shift}:B{2 + shift}"
    assert main.freeze_panes == f"B{2 + shift}"
    assert main.row_breaks.brk[0].id == 2 + shift
    assert main.print_title_cols == "$A:$A"
    assert workbook_identities(path) == {"Candidate": identity}
    with deannotated_workbook(path) as restored:
        assert restored.read_bytes() == original
    assert path.read_bytes() != original
    annotate_workbook(path, {"Candidate": identity})
    assert workbook_identities(path) == {"Candidate": identity}
    with deannotated_workbook(path) as restored:
        assert restored.read_bytes() == original


def test_real_dependency_dashboards_are_readable_annotated_and_historical(tmp_path: Path) -> None:
    from digital_metrics.reporting.dashboard import get_dashboards
    from digital_metrics.types import Metrics

    get_dashboards(
        {"cat": Metrics(tp=4, fp=1, fn=2, confidence=0.377)},
        pd.DataFrame({"instance_label": ["cat"], "split": ["test"]}),
        None,
        ["cat"],
        path=str(tmp_path),
    )
    for path in tmp_path.glob("*.xlsx"):
        before = pd.read_excel(path, index_col=0)
        assert workbook_identities(path) == {}
        pd.testing.assert_frame_equal(read_dashboard(path, index_col=0), before)
        annotate_workbook(
            path, {"Model": ModelIdentity(model_name="+safe", training_task_id="task")}
        )
        pd.testing.assert_frame_equal(read_dashboard(path, index_col=0), before)
        with ZipFile(path) as archive:
            assert b"+safe" in archive.read("xl/worksheets/sheet1.xml")


def test_real_report_workbooks_keep_every_original_sheet_and_cell(tmp_path: Path) -> None:
    from digital_metrics.reporting.dashboard import get_dashboards
    from digital_metrics.types import Metrics
    from report_generator.config import Config
    from report_generator.core.reader import MetricsReader
    from report_generator.reports.business.builder import BusinessReportBuilder
    from report_generator.reports.dev.builder import DevReportBuilder

    get_dashboards(
        {"cat": Metrics(tp=4, fp=1, fn=2, confidence=0.37)},
        pd.DataFrame({"instance_label": ["cat"], "split": ["test"]}),
        None,
        ["cat"],
        path=str(tmp_path),
    )
    dashboard = tmp_path / "full_dashboard_default.xlsx"
    baseline = ModelIdentity(model_name="Old model", training_task_id="old-task")
    candidate = ModelIdentity(model_name="新模型", training_task_id="new-task")
    reader = MetricsReader(dashboard)
    expected = reader.read()
    annotate_workbook(dashboard, {"Candidate": candidate})
    with deannotated_workbook(dashboard) as clean:
        adapted = MetricsReader(clean)
        pd.testing.assert_frame_equal(adapted.read(), expected)
        dev = tmp_path / "dev.xlsx"
        business = tmp_path / "business.xlsx"
        DevReportBuilder(adapted, adapted, Config.load()).build(dev)
        BusinessReportBuilder(adapted, adapted, Config.load()).build(business)
    for path in (dev, business):
        before = load_workbook(path)
        annotate_workbook(path, {"Baseline": baseline, "Candidate": candidate})
        after = load_workbook(path)
        assert before.sheetnames == after.sheetnames
        for old, new in zip(before, after, strict=True):
            assert new["A1"].value == "Baseline: Old model"
            assert new["A3"].value == "Candidate: 新模型"
            for row in old:
                for cell in row:
                    if cell.data_type != "f":
                        assert new.cell(cell.row + 4, cell.column).value == cell.value
            assert new.print_title_rows.startswith("$1:$4")


def test_chart_validation_and_filter_ranges_follow_original_body(tmp_path: Path) -> None:
    from openpyxl.chart import BarChart, Reference  # type: ignore[import-untyped]
    from openpyxl.worksheet.datavalidation import DataValidation  # type: ignore[import-untyped]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Values, original"
    sheet.append(["Class", "Value"])
    sheet.append(["cat", 4])
    sheet.auto_filter.ref = "A1:B2"
    validation = DataValidation(type="whole", formula1="B2", formula2="10")
    validation.add("B2")
    sheet.add_data_validation(validation)
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    sheet.add_chart(chart, "D2")
    sheet.print_title_rows = "1:1"
    sheet.print_title_cols = "A:A"
    sheet.page_setup.fitToHeight = 2
    path = tmp_path / "chart.xlsx"
    workbook.save(path)
    annotate_workbook(path, {"Model": ModelIdentity(model_name="model", training_task_id="task")})
    after = load_workbook(path)["Values, original"]
    assert after.auto_filter.ref == "A3:B4"
    assert after.data_validations.dataValidation[0].sqref == "B4"
    assert after.data_validations.dataValidation[0].formula1 == "B4"
    assert after.print_title_rows == "$1:$3"
    assert after.print_title_cols == "$A:$A"
    assert after.page_setup.fitToHeight == 2
    with ZipFile(path) as archive:
        drawing = archive.read("xl/drawings/drawing1.xml")
        assert b">3</" in drawing  # chart's zero-based D2 anchor now follows D4
        chart_xml = archive.read("xl/charts/chart1.xml")
        assert b"'Values, original'!$B$4" in chart_xml


def test_edited_annotated_parts_are_rejected_before_reading_original(tmp_path: Path) -> None:
    import io

    import pytest

    path = tmp_path / "edited.xlsx"
    pd.DataFrame({"class": ["cat"], "confidence": [0.37]}).to_excel(path, index=False)
    annotate_workbook(path, {"Model": ModelIdentity(model_name="model", training_task_id="task")})
    content = path.read_bytes()
    with ZipFile(io.BytesIO(content)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    worksheet = "xl/worksheets/sheet1.xml"
    assert b">0.37<" in parts[worksheet]
    parts[worksheet] = parts[worksheet].replace(b">0.37<", b">0.91<")
    with ZipFile(path, "w") as archive:
        for name, part in parts.items():
            archive.writestr(name, part)
    with pytest.raises(ValueError, match="changed after model identity annotation"):
        workbook_identities(path)
    with pytest.raises(ValueError, match="changed after model identity annotation"):
        read_dashboard(path)
    with (
        pytest.raises(ValueError, match="changed after model identity annotation"),
        deannotated_workbook(path),
    ):
        pytest.fail("Edited workbook must not yield stale original content")


def test_malformed_model_metadata_is_rejected_by_all_read_adapters(tmp_path: Path) -> None:
    import io

    import pytest

    path = tmp_path / "malformed.xlsx"
    pd.DataFrame({"class": ["cat"], "confidence": [0.37]}).to_excel(path, index=False)
    annotate_workbook(path, {"Model": ModelIdentity(model_name="model", training_task_id="task")})
    with ZipFile(io.BytesIO(path.read_bytes())) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    metadata = "customXml/clearml-yolo-identity.xml"
    parts[metadata] = parts[metadata].replace(b'"model_name": "model"', b'"model_name": ""')
    with ZipFile(path, "w") as archive:
        for name, part in parts.items():
            archive.writestr(name, part)
    with pytest.raises(ValueError, match="Invalid workbook model identity metadata"):
        workbook_identities(path)
    with pytest.raises(ValueError, match="Invalid workbook model identity metadata"):
        read_dashboard(path)
    with (
        pytest.raises(ValueError, match="Invalid workbook model identity metadata"),
        deannotated_workbook(path),
    ):
        pytest.fail("Malformed identities must not yield original content")


@pytest.mark.parametrize("changed_part", ["identity", "original"])
def test_valid_metadata_edits_are_rejected(tmp_path: Path, changed_part: str) -> None:
    import base64
    import io

    path = tmp_path / "edited-metadata.xlsx"
    pd.DataFrame({"class": ["cat"], "confidence": [0.37]}).to_excel(path, index=False)
    original = path.read_bytes()
    annotate_workbook(path, {"Model": ModelIdentity(model_name="model", training_task_id="task")})
    with ZipFile(io.BytesIO(path.read_bytes())) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    metadata = "customXml/clearml-yolo-identity.xml"
    if changed_part == "identity":
        parts[metadata] = parts[metadata].replace(
            b'"training_task_id": "task"', b'"training_task_id": "other"'
        )
    else:
        edited = io.BytesIO()
        with ZipFile(io.BytesIO(original)) as source, ZipFile(edited, "w") as destination:
            for name in source.namelist():
                content = source.read(name)
                if name == "xl/worksheets/sheet1.xml":
                    assert b">0.37<" in content
                    content = content.replace(b">0.37<", b">0.91<")
                destination.writestr(name, content)
        parts[metadata] = parts[metadata].replace(
            base64.b64encode(original), base64.b64encode(edited.getvalue())
        )
    with ZipFile(path, "w") as archive:
        for name, part in parts.items():
            archive.writestr(name, part)
    with pytest.raises(ValueError, match="changed after model identity annotation"):
        workbook_identities(path)
    with pytest.raises(ValueError, match="changed after model identity annotation"):
        read_dashboard(path)
